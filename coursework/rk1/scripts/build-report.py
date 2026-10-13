"""Build the submission report from real recorded lab evidence (python-docx)."""
from pathlib import Path
import json
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
OUT = Path(sys.argv[1])
OUT.mkdir(parents=True, exist_ok=True)
lab = json.loads((ROOT / 'evidence/localnet-run.json').read_text())
mvp = json.loads((ROOT / 'evidence/mvp-integration.json').read_text())
detective = json.loads((ROOT / 'evidence/devnet-detective.json').read_text())
tx = {t['label']: t for t in lab['transactions']}

doc = Document()
for border in list(doc.styles.element.xpath('.//w:pBdr')):
    border.getparent().remove(border)
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(.78)
section.bottom_margin = Inches(.75)
section.left_margin = section.right_margin = Inches(.85)
section.header_distance = section.footer_distance = Inches(.32)
section.different_first_page_header_footer = True
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.get_or_add_rPr().append(OxmlElement('w:lang'))
    style._element.rPr[-1].set(qn('w:val'), 'ru-RU')
normal = doc.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.12
normal.paragraph_format.space_after = Pt(7)
for name, size in [('Title', 27), ('Heading 1', 19), ('Heading 2', 13)]:
    st = doc.styles[name]
    st.font.size = Pt(size)
    st.paragraph_format.space_before = Pt(12 if name == 'Heading 2' else 0)
    st.paragraph_format.space_after = Pt(9)
    st.paragraph_format.keep_with_next = True
doc.styles['Subtitle'].font.size = Pt(14)
doc.styles['Title'].font.bold = True
header = section.header.paragraphs[0]
header.text = 'EVENTLY   ·   ПРАКТИЧЕСКИЕ РАБОТЫ И РК1'
header.runs[0].font.size = Pt(9)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Evently  |  ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
for run in footer.runs:
    run.font.size = Pt(9)
doc.core_properties.title = 'Evently Практические работы и РК1'
doc.core_properties.subject = 'Работы по силлабусу за недели 1–8'
doc.core_properties.author = 'Кочетов Руслан'
doc.core_properties.keywords = 'Solana, Rust, Anchor, PDA, Token-2022, Evently, РК1'
md = []

def p(text='', bold=False):
    para = doc.add_paragraph()
    run = para.add_run(text); run.bold = bold
    md.append(text + '\n')
    return para

def heading(text, level=1):
    doc.add_heading(text, level)
    md.append('#' * (level + 1) + ' ' + text + '\n')

def page(title):
    doc.add_page_break(); md.append('\n<!-- pagebreak -->\n')
    heading(title)

def code(text):
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(3)
    para.paragraph_format.space_after = Pt(9)
    para.paragraph_format.line_spacing = 1
    para.paragraph_format.keep_together = True
    run = para.add_run(text)
    run.font.name = 'DejaVu Sans Mono'; run.font.size = Pt(9)
    md.append('```\n' + text + '\n```\n')

def link(label, url):
    para = doc.add_paragraph()
    rel = para.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement('w:hyperlink'); h.set(qn('r:id'), rel)
    r = OxmlElement('w:r'); props = OxmlElement('w:rPr')
    c = OxmlElement('w:color'); c.set(qn('w:val'), '185A8D'); props.append(c)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); props.append(u)
    r.append(props); t = OxmlElement('w:t'); t.text = label; r.append(t); h.append(r); para._p.append(h)
    md.append(f'[{label}]({url})\n')

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths): col.width = Inches(width)
    for i, name in enumerate(headers): t.rows[0].cells[i].text = name
    repeat = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row): cells[i].text = str(value)
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        e = OxmlElement('w:' + edge)
        for k, v in {'val':'single', 'sz':'4', 'color':'D9D9D9'}.items(): e.set(qn('w:' + k), v)
        borders.append(e)
    t._tbl.tblPr.append(borders)
    for ri, row in enumerate(t.rows):
        cant = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(cant)
        for ci, cell in enumerate(row.cells):
            cell.width = Inches(widths[ci]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for edge in ['top','bottom','left','right']:
                e = OxmlElement('w:' + edge); e.set(qn('w:w'), '85'); e.set(qn('w:type'),'dxa'); margins.append(e)
            pr.append(margins)
            if ri == 0:
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E3EBF1'); pr.append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(2)
                para.paragraph_format.space_before = Pt(2)
                para.paragraph_format.line_spacing = 1.05
                if ci == 0 and widths[ci] < 1: para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    run.font.size = Pt(10.5); run.bold = ri == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    md.append('| ' + ' | '.join(headers) + ' |\n|' + '|'.join(['---']*len(headers)) + '|\n' + '\n'.join('| ' + ' | '.join(str(x).replace('\n','; ') for x in row) + ' |' for row in rows) + '\n')

def evidence(text):
    para = p('Доказательства в архиве: ' + text)
    for run in para.runs: run.font.size = Pt(10)

# Cover
p('ALMATY MANAGEMENT UNIVERSITY', True).alignment = WD_ALIGN_PARAGRAPH.CENTER
p('School of Digital Technologies and Economics').alignment = WD_ALIGN_PARAGRAPH.CENTER
p('ГОП B057 ИТ  ·  Семестр 1').alignment = WD_ALIGN_PARAGRAPH.CENTER
p('\n')
doc.add_paragraph('Evently', 'Title')
doc.add_paragraph('Практические работы и РК1', 'Title')
p('Отчёт по заданиям недель 1–8')
p('Дисциплина «ИТ-стартап Итоговый проект 1 (DS/ML стартап)»')
p('Платформа мероприятий с проверяемыми билетами в Solana')
p('\n')
p('Студент  Кочетов Руслан')
p('Группа    Software Engineering')
p('Язык обучения  Русский')
p('Преподаватель  Иманкулова Алия')
p('Дата проверки проекта  4 октября 2026 года')
p('\n')
p('В комплект входят исходный код, результаты запуска и инструкция воспроизведения. '
  'Локальные проверки успешно выполнены. Публикация собственных программ в Devnet '
  'ожидает пополнения учебного кошелька тестовыми SOL; это ограничение отражено в отчёте.')
p('Алматы 2026').alignment = WD_ALIGN_PARAGRAPH.CENTER
md.insert(0, '# Evently Практические работы и РК1\n')

page('Объём работы и соответствие силлабусу')
p('Отчёт охватывает семь практических работ и первый рубежный контроль из силлабуса, '
  'утверждённого 19 августа 2026 года. Задания реализованы на основе Evently. '
  'Работы по DeFi, CPI и финальному Capstone из недель 9–15 в этот комплект не входят.')
table(['Неделя', 'Задание и вес по силлабусу', 'Фактический результат'], [
 ('1','Кошелёк и первая транзакция · 2,5%','Перевод SOL подтверждён локально; собственная Devnet-транзакция ожидает SOL'),
 ('2','Transaction Detective · 2,5%','Разобраны 9 локальных транзакций и 1 публичная транзакция Devnet'),
 ('3','Rust CLI · 5%','Программа регистрации и check-in; 4 теста пройдены'),
 ('4','Первая Solana Program · 5%','Собрана и развёрнута локально; требование Devnet пока не закрыто'),
 ('5','Anchor Counter · 5%','Создание и изменение счётчика через клиент'),
 ('6','Профиль PDA · 5%','Создание, обновление и защита authority'),
 ('7','Свой токен · 5%','Выпуск, metadata, transfer и чтение состояния Token-2022'),
 ('8','РК1 Web3 MVP · 5%','Сквозной сценарий Evently с реальным локальным блокчейном'),
], [.65, 2.65, 3.5])
p('Указанные проценты — веса заданий из силлабуса, а не выставленные оценки. Их сумма за недели 1–8 составляет 35%.')
heading('Среда проверки', 2)
p('macOS, Docker Desktop, Rails 8.1/PostgreSQL, React 19/Vinext, Node.js, Rust, '
  'Agave CLI 2.3.0, Anchor 0.32.1. Учебные программы используют platform-tools v1.51; '
  'совместимые зависимости зафиксированы в Cargo.lock и package-lock.json.')
p('Localnet работает на 127.0.0.1:8899. Учебный API использует отдельную базу evently_rk1 '
  'и порт 8001. Основная база Evently не подменялась учебными результатами.')

page('Неделя 1 Кошелёк и первая транзакция')
p('Цель — создать кошелёк, получить тестовые SOL и проследить перевод от подписи до подтверждения. '
  'Создан отдельный учебный keypair. Приватные ключи пользователя и Phantom в этой работе не используются.')
heading('Ход выполнения', 2)
p('Локальный валидатор выдал учебному кошельку 100 SOL. Клиент сформировал инструкцию '
  'SystemProgram.transfer на 1 000 000 lamports, подписал транзакцию и дождался статуса confirmed. '
  'RPC-ответ сохранён вместе с accounts, logs, slot и балансами.')
p('Адрес отправителя:'); code(lab['payer'])
p('Адрес получателя:'); code(lab['recipient'])
table(['Показатель', 'Результат'], [
 ('Кластер','Localnet'), ('Сумма перевода','0,001 SOL = 1 000 000 lamports'),
 ('Комиссия',f"{tx['week1_sol_transfer']['fee_lamports']} lamports = 0,000005 SOL"),
 ('Slot',tx['week1_sol_transfer']['slot']), ('Ошибка транзакции','meta.err = null'),
 ('Изменение баланса отправителя','−1 005 000 lamports'),
 ('Изменение баланса получателя','+1 000 000 lamports'),
], [3.1,3.7])
heading('Разбор результата', 2)
p('Отправитель является signer и writable account; получатель — writable account. '
  'System Program исполняет перевод. Комиссия списывается дополнительно к сумме перевода; '
  '1 SOL равен 10⁹ lamports [2].')
p('Собственная публичная транзакция Devnet пока не отправлена: баланс учебного адреса равен нулю, '
  'автоматический faucet вернул ограничение. Локальную signature нельзя выдавать за запись в Devnet Explorer. '
  'После пополнения повторный запуск создаст действительную ссылку Explorer.')
evidence('coursework/rk1/evidence/localnet-run.json, запись week1_sol_transfer; wallet.json; client/run.mjs.')

page('Неделя 2 Разбор реальных транзакций')
p('Цель — определить accounts, instructions и programs, а затем объяснить изменение состояния. '
  'Выполнен разбор собственных локальных транзакций и отдельного публичного примера Devnet. '
  'Публичный пример отправлен сторонним пользователем и не приписывается автору проекта.')
link('Открыть исследованную Devnet-транзакцию в Solana Explorer', detective['explorer'])
p(f"RPC подтвердил успешное выполнение в slot {detective['slot']}. Транзакция содержит "
  f"{len(detective['accounts'])} accounts, {len(detective['instructions'])} внешние инструкции и комиссию "
  f"{detective['fee_lamports']} lamports. Первый account — плательщик комиссии и единственный внешний signer.")
table(['Часть транзакции', 'Наблюдение и смысл'], [
 ('Инструкции 1–3','Associated Token Account Program: createIdempotent. Вызовы обеспечивают наличие token accounts; повторное создание уже существующего ATA не требуется.'),
 ('Инструкция 4','Пользовательская программа XuXkWA…ZMoRfi. В логе указано Instruction: Buy; без IDL её параметры полностью не декодируются.'),
 ('Вложенные вызовы','Три transfer и один mintTo в SPL Token Program, затем вложенный вызов пользовательской программы.'),
 ('Lamports','У плательщика списано 5000; остальные lamport-балансы из ответа не изменились. Токены учитываются отдельно от SOL.'),
], [1.6,5.2])
p('Вложенные transfer содержат amounts 14 324 001, 70 562 и 466 866; mintTo — 466 866. '
  'Это целые значения минимальных единиц. Без decimals соответствующих mint их нельзя корректно '
  'называть количеством отображаемых токенов. По одному имени Buy нельзя доказывать экономический смысл операции.')
heading('Сравнение со своим переводом', 2)
p('Локальный SOL transfer недели 1 содержит одну инструкцию System Program, двух участников '
  'и простую проверяемую разницу балансов. Создание токена недели 7 включает несколько программ '
  'и rent-exempt deposit. Разбор отделяет комиссию от средств, внесённых для хранения аккаунтов [2, 3].')
evidence('evidence/devnet-detective.json; localnet-transaction-analysis.json; client/detective.mjs и client/analyze.mjs.')

page('Неделя 3 Программа на Rust')
p('Реализована CLI-программа evently-cli, моделирующая регистрацию на событие ограниченной вместимости '
  'и однократный проход по билету. Это небольшая самостоятельная программа, которую можно запустить без блокчейна.')
heading('Модель и конструкции языка', 2)
p('Структура Event хранит название, capacity и Vec<Ticket>. Ticket содержит owner и статус Active или Used. '
  'Ошибки представлены enum BookingError: SoldOut, Duplicate, NotFound и AlreadyUsed. '
  'Методы возвращают Result, а изменение состояния требует &mut self. Поиск использует итераторы, '
  'а данные события сериализуются в JSON.')
code('cargo run --manifest-path coursework/rk1/Cargo.toml \\\n  -p evently-cli -- demo 2\n\ncargo test --manifest-path coursework/rk1/Cargo.toml \\\n  -p evently-cli')
heading('Проверка правил', 2)
table(['Тест', 'Ожидаемый результат', 'Факт'], [
 ('capacity_is_enforced','Второй гость при capacity = 1 получает SoldOut','Пройден'),
 ('duplicate_does_not_consume_seat','Повторная регистрация не увеличивает число билетов','Пройден'),
 ('check_in_is_single_use','Повторный проход получает AlreadyUsed','Пройден'),
 ('unknown_guest_is_rejected','Неизвестный гость получает NotFound','Пройден'),
], [2.4,3.6,.8])
p('В демонстрации guest-alice получает билет и проходит один раз. Итоговый JSON содержит статус Used; '
  'повторный вызов выводит Repeated check-in: AlreadyUsed. Аргумент вместимости ограничен диапазоном 1–5000.')
p('Связь с Solana: явные структуры состояния, enum ошибок, заимствования и проверка переходов '
  'используются затем в on-chain программе. CLI не хранит данные между запусками и не заменяет серверную базу Evently.')
evidence('cli/src/main.rs; evidence/rust-tests.txt; evidence/rust-cli.txt.')

page('Неделя 4 Первая Solana Program')
p('Создана нативная программа hello_evently на Rust с entrypoint. Она получает список аккаунтов '
  'и байты инструкции, требует подпись первого аккаунта и принимает только строку evently-rk1. '
  'При успехе программа записывает сообщение Hello Evently RK1 и адрес signer в журнал исполнения.')
p('Program ID в проверенном локальном развёртывании:'); code(lab['hello_program'])
heading('Сборка и проверка', 2)
code('cargo build-sbf \\\n  --manifest-path programs/hello_evently/Cargo.toml \\\n  --sbf-out-dir target/deploy --tools-version v1.51\n\nbash scripts/deploy.sh localnet\nnode client/run.mjs localnet')
p('Получен исполняемый SBF-файл hello_evently.so. Программа развёрнута на локальном валидаторе; '
  'клиент проверил executable account и выполнил инструкцию. Успешный вызов подтверждён в '
  f"slot {tx['week4_hello_program']['slot']}; комиссия составила 5000 lamports.")
table(['Проверка', 'Результат'], [
 ('Подпись и правильные байты','Успешное выполнение, meta.err = null'),
 ('Инструкция wrong','Отклонена с InvalidInstructionData'),
 ('Отсутствие signer','В исходном коде возвращается MissingRequiredSignature; отдельный runtime-тест не заявляется'),
], [2.7,4.1])
heading('Статус требования Devnet', 2)
p('Силлабус прямо требует развёртывание в Devnet. В данном комплекте подтверждено локальное '
  'развёртывание; публичное ещё не выполнено из-за отсутствия тестовых SOL. Это незакрытая часть '
  'работы недели 4, а не эквивалентный результат. Скрипт поддерживает Devnet и готов к запуску после пополнения.')
p('У программы нет собственного хранилища состояния. Состояние добавляется в работах 5–6 '
  'через отдельные accounts, которыми владеет Anchor-программа [3].')
evidence('programs/hello_evently/src/lib.rs; evidence/localnet-deploy-hello_evently.json; localnet-run.json.')

page('Неделя 5 Счётчик на Anchor')
p('Anchor-программа evently_lab реализует initialize_counter и increment. Для каждого authority '
  'счётчик имеет отдельный PDA с seeds ["counter", authority]. Клиент вычисляет адрес, '
  'инициализирует аккаунт, читает значение, отправляет increment и читает состояние повторно [4].')
p('Program ID:'); code(lab['anchor_program'])
p('Counter PDA:'); code(lab['counter']['address'])
table(['Поле', 'Тип', 'Проверенное значение'], [
 ('authority','Pubkey','Учебный кошелёк отправителя'),
 ('count','u64',f"{lab['counter']['before']} → {lab['counter']['after']}"),
 ('bump','u8',lab['counter']['bump']),
 ('Размер','bytes','49 = 8 discriminator + 32 authority + 8 count + 1 bump'),
], [1.6,1.2,4.0])
heading('Контроль доступа и переполнения', 2)
code('#[account(mut,\n  seeds = [b"counter", authority.key().as_ref()],\n  bump = counter.bump, has_one = authority)]\npub counter: Account<\'info, Counter>,\npub authority: Signer<\'info>,')
p('Signer требует подпись authority. Seeds и bump связывают переданный адрес с ожидаемым PDA, '
  'а has_one проверяет владельца права изменения в сохранённом состоянии. '
  'Инкремент использует checked_add и возвращает Overflow вместо переполнения.')
p('Положительная проверка подтвердила увеличение ровно на единицу. Вызов от второго кошелька '
  'с тем же account отклонён ограничением Anchor. IDL сгенерирован из программы и используется клиентом; '
  'неуспешный вызов не принимается как успешная транзакция.')
evidence('programs/evently_lab/src/lib.rs; idl/evently_lab.json; client/run.mjs; localnet-run.json, раздел counter.')

page('Неделя 6 Профиль с использованием PDA')
p('Реализован on-chain профиль организатора с инструкциями create_profile и update_profile. '
  'Профиль выводится из seeds ["profile", authority] и Program ID. Для одного authority '
  'в этой программе получается один детерминированный адрес; приватного ключа у PDA нет [4].')
p('Profile PDA:'); code(lab['profile']['address'])
table(['Параметр', 'Результат'], [
 ('Bump',lab['profile']['bump']),
 ('Name',lab['profile']['data']['name']),
 ('Bio',lab['profile']['data']['bio']),
 ('Owner account','Anchor-программа evently_lab'),
 ('Authority','Учебный кошелёк, подписывающий изменения'),
 ('Максимальный размер','241 bytes с discriminator и префиксами длины строк'),
], [2.1,4.7])
heading('Проверки состояния', 2)
p('Клиент создал профиль, обновил описание, прочитал account и сравнил сохранённые значения с '
  'ожидаемыми. Отдельно проверены owner программы и authority кошелька. Это разные понятия: '
  'owner определяет программу, управляющую данными account; authority является полем модели доступа.')
p('Изменение от чужого кошелька отклонено. Имя из 33 ASCII-символов отклонено с InvalidName. '
  'Имя ограничено 32 байтами UTF-8 и не может быть пустым после trim; описание ограничено 160 байтами. '
  'Для кириллицы число допустимых символов меньше числа байтов. Отдельный тест границы bio не заявляется.')
heading('Приватность и применение', 2)
p('В блокчейн записаны только учебный псевдоним и публичное описание. Email, пароль и анкета гостя '
  'остаются вне этого PDA. Профиль является учебным расширением: существующий интерфейс Evently '
  'пока не читает его автоматически. Доступность данных через RPC не означает их конфиденциальность.')
evidence('programs/evently_lab/src/lib.rs; localnet-run.json, раздел profile и проверки profile_rejects_*; IDL.')

page('Неделя 7 Выпуск и перевод токена')
p('Создан учебный токен Evently Lab Credit с символом ELAB. Он использует Token-2022, '
  'расширение MetadataPointer и встроенные TokenMetadata. Метаданные содержат название, символ '
  'и URI JSON; клиент читает их обратно из mint account [5].')
p('Mint:'); code(lab['token']['mint'])
table(['Показатель', 'Значение'], [
 ('Decimals','2'), ('Выпущено','100 000 минимальных единиц = 1000 ELAB'),
 ('Переведено второму кошельку','2500 минимальных единиц = 25 ELAB'),
 ('Остаток у отправителя','97 500 минимальных единиц = 975 ELAB'),
 ('Баланс получателя','2500 минимальных единиц = 25 ELAB'),
 ('Mint authority после выпуска','null — право дополнительной эмиссии отозвано'),
 ('Freeze authority','null'),
], [3.1,3.7])
heading('Последовательность операций', 2)
p('Первой транзакцией создаётся mint account, инициализируются MetadataPointer, mint и metadata. '
  'Второй создаются ATA отправителя и получателя, выпускаются токены и отзывается mint authority. '
  'Третья выполняет TransferChecked: decimals проверяются вместе с суммой. Клиент затем сверяет '
  'оба баланса, supply, отсутствие mint authority и название metadata.')
p('URI метаданных указывает на public/metadata/evently-lab.json в репозитории. Этот JSON также '
  'включён в архив; при публикации своей копии проекта нужно использовать собственный доступный URI. '
  'Metadata update authority сохраняется у учебного кошелька, поэтому метаданные не заявляются неизменяемыми.')
heading('Связь с билетами Evently', 2)
p('ELAB создан отдельно для задания transfer и не используется как платёжное средство. '
  'Билет Evently — другой mint: decimals 0, supply 1, NonTransferable и отозванное право эмиссии. '
  'Передаваемость учебного токена не меняет правило непередаваемости билета [6].')
evidence('client/run.mjs; localnet-run.json, раздел token; записи week7_*; public/metadata/evently-lab.json.')

page('РК1 Концепция и архитектура MVP')
p('Evently решает задачу организации мероприятий: публикация события, регистрация гостя, '
  'выдача билета и контроль однократного входа. Web3-компонент позволяет независимо проверить '
  'факт выпуска токена и его связь с кошельком. Проверка вместимости и допуска остаётся в серверной системе.')
heading('Разделение ответственности', 2)
table(['Компонент', 'Ответственность'], [
 ('React 19 и Vinext','Каталог, регистрация аккаунта, подключение кошелька, билеты, QR и кабинет организатора'),
 ('Ruby on Rails','Авторизация, роли и ownership, резерв мест, очередь, проверка Solana и атомарный check-in'),
 ('PostgreSQL','Пользователи, события, регистрации, билеты и использованные коды'),
 ('Phantom и Solana','Подпись challenge и транзакций; Token-2022 mint, token account и memo'),
], [2.0,4.8])
heading('Основной сценарий', 2)
p('Гость входит по email и паролю, затем подтверждает владение кошельком подписью одноразового '
  'challenge. Сервер резервирует место. Кошелёк подписывает выпуск непередаваемого Token-2022 '
  'и memo с идентификаторами события и регистрации. Rails проверяет подтверждённую транзакцию '
  'через RPC и только затем помечает билет confirmed.')
p('На входе организатор подписывает memo, относящееся к конкретному билету. Сервер проверяет '
  'подпись, принадлежность токена гостю и права организатора; в транзакции базы данных '
  'меняет билет active → used. Второй проход блокируется. Сам токен не кодирует окончательное '
  'решение о допуске: его принимает Evently на основании серверного состояния.')
heading('Границы прототипа', 2)
p('Роли прототипа: visitor, organizer, admin. Подтверждение email, реальные платежи и SSO '
  'не входят в проверенную реализацию РК1. Counter и Profile PDA демонстрируют отдельные '
  'учебные программы; основной билетный поток использует стандартный Token-2022. '
  'Факт MVP не означает готовность всех функций ранее обсуждённой продуктовой спецификации.')
link('Репозиторий Evently', 'https://github.com/miracleunvrs/evently')

page('РК1 Проверки сквозного сценария')
p('Проверка выполнялась 4 октября 2026 года в изолированной базе evently_rk1. '
  'Клиент использовал настоящие Ed25519-подписи учебных кошельков и отправлял транзакции '
  'локальному Solana-валидатору. Ответы Rails и состояния билета проверялись утверждениями assert.')
table(['Проверяемое действие', 'Полученный результат'], [
 ('Вход visitor, organizer, admin','HTTP 200; роли соответствуют назначенным'),
 ('Подтверждение кошельков','Challenge и проверка подписей: HTTP 200'),
 ('Создание события организатором','HTTP 201; вместимость 1'),
 ('Регистрация гостя и проверка Token-2022','HTTP 201 и 200; билет confirmed'),
 ('Чтение списка гостей посетителем','HTTP 404'),
 ('Создание события посетителем','HTTP 403'),
 ('Очередь при заполнении события','HTTP 201'),
 ('Первый check-in','HTTP 200'),
 ('Повторный check-in','HTTP 409'),
 ('Отмена использованного билета','HTTP 409'),
 ('Итоговая статистика','1 регистрация, 1 посещение, capacity 1'),
], [3.5,3.3])
p(f"В успешном сохранённом запуске выполнены {len(mvp['checks'])} проверки. Сценарий использовал уже созданные "
  'учебные аккаунты; он не выдаёт повторный вход за повторную регистрацию. При первом запуске '
  'клиент создаёт аккаунты через публичную форму API, а учебные роли назначает только в изолированной development-базе.')
p('Дополнительно пройдены 22 Rails-теста с 52 assertions, 8 frontend-тестов и 4 Rust-теста. '
  'Лабораторный клиент записал 9 подтверждённых транзакций и 5 проверок отклонений или итогового состояния токена. '
  'Эти числа описывают выполненные проверки, а не гарантируют отсутствие всех ошибок.')
evidence('evidence/mvp-integration.json; rails-tests.txt; frontend-tests.txt; rust-tests.txt; localnet-run.json.')

page('Воспроизведение и завершение Devnet')
heading('Запуск из архива', 2)
p('Распакуйте каталог evently, установите Node.js, Docker, Rust, Solana CLI 2.3.0 и Anchor CLI 0.32.1. '
  'Полные команды находятся в coursework/rk1/README.md. Сначала установите зависимости, '
  'создайте собственные учебные ключи и синхронизируйте Program ID, затем соберите программы и IDL.')
code('npm ci\nnpm ci --prefix coursework/rk1\nnode coursework/rk1/scripts/wallet.mjs\nnode coursework/rk1/scripts/sync-program-id.mjs')
p('Запустите локальный валидатор, выдайте SOL обоим учебным кошелькам и выполните deploy.sh localnet. '
  'Клиенты run.mjs и analyze.mjs пересоздадут доказательства лабораторных работ. '
  'Для MVP поднимите отдельный API на порту 8001 и выполните client/mvp.mjs. '
  'Воспроизводимость не требует раскрытия исходных приватных ключей: новые адреса и signatures будут отличаться.')
heading('Что осталось для полного соответствия', 2)
p('На момент подготовки отчёта учебный адрес в Devnet имеет нулевой баланс. '
  'Автоматический запрос тестовых SOL ограничен faucet. После получения тестовых SOL '
  'нужно выполнить команды ниже из coursework/rk1 и приложить новые JSON и ссылки Explorer. '
  'Покупка SOL и использование Mainnet не требуются.')
code('bash scripts/deploy.sh devnet\nnode client/run.mjs devnet\nnode client/analyze.mjs devnet')
p('Успех подтверждают оба файла devnet-deploy-*.json, файл devnet-run.json с completed_at '
  'и успешными проверками, а также реальные Devnet Explorer-ссылки. '
  'Пока эти файлы не получены, собственный Devnet-перевод недели 1 и Devnet-деплой недели 4 не считаются выполненными.')
heading('Состав и безопасность архива', 2)
p('Архив содержит исходники Evently, учебные программы, IDL, lockfiles, доказательства запусков '
  'и инструкцию защиты. Приватные keypair, пароли, .env, зависимости, дампы базы и локальный ledger исключены. '
  'ФИО и группа на титульном листе заполняются перед отправкой преподавателю.')

page('План демонстрации и вопросы защиты')
p('Предлагаемая демонстрация занимает около семи минут. Это рабочий сценарий подготовки; '
  'силлабус не задаёт отдельную длительность защиты РК1.')
table(['Время', 'Демонстрация'], [
 ('0–1 мин','Проблема Evently, ценность проверяемого билета и разделение базы данных и блокчейна'),
 ('1–2 мин','Rust CLI и тесты регистрации и повторного входа'),
 ('2–3 мин','Native программа, счётчик Anchor и изменение значения'),
 ('3–4 мин','Profile PDA: seeds, bump, owner и отклонение чужой подписи'),
 ('4–5 мин','Транзакции, ELAB, metadata и перевод 25 токенов'),
 ('5–7 мин','Билет Evently, check-in, повторный проход и ограничения прототипа'),
], [1.0,5.8])
heading('Ключевые ответы', 2)
p('Почему нужен PDA? Он даёт воспроизводимый адрес состояния по seeds и Program ID. '
  'В профиле право изменения дополнительно проверяется по подписи authority; знание адреса не даёт права записи.')
p('Что обеспечивает single-use? Атомарный переход active → used в Rails/PostgreSQL. '
  'Memo фиксирует подписанное действие в блокчейне, но сам по себе не заменяет серверную проверку повторного входа.')
p('Почему токен ELAB переводится, а билет нет? Это разные mint accounts. ELAB нужен для '
  'учебного transfer; у билета активировано расширение NonTransferable и supply равен единице.')
p('Что даёт Anchor? Проверки accounts и signer, сериализацию, discriminator и IDL. '
  'Правила бизнеса, размеры данных и корректные constraints остаются ответственностью разработчика.')
p('Что нельзя утверждать по этому отчёту? Что локальные signatures существуют в Devnet, '
  'что все планы продукта уже реализованы, что передача токена автоматически разрешает вход или '
  'что прохождение тестов доказывает отсутствие всех уязвимостей.')
evidence('Подробный сценарий и дополнительные вопросы: coursework/rk1/docs/defense.md.')

page('Источники и реестр доказательств')
p('Требования к заданиям взяты из предоставленного силлабуса. Технические определения сверены '
  'с официальной документацией Solana и Anchor. Результаты исполнения взяты из файлов evidence; '
  'дата фиксации — 4 октября 2026 года.')
p('[1] Силлабус «ИТ-стартап Итоговый проект 1 (DS/ML стартап)». AlmaU. '
  'Иманкулова А.Ж. Утверждён 19.08.2026. Раздел недель 1–8 и критерии оценивания.')
sources = [
 ('[2] Solana Transactions', 'https://solana.com/docs/core/transactions'),
 ('[3] Solana Accounts', 'https://solana.com/docs/core/accounts'),
 ('[4] Anchor Program Derived Addresses', 'https://www.anchor-lang.com/docs/basics/pda'),
 ('[5] Solana Token Metadata', 'https://solana.com/docs/tokens/extensions/metadata'),
 ('[6] Solana Token Extensions', 'https://solana.com/docs/tokens/extensions'),
 ('[7] Anchor 0.32.1 Release Notes', 'https://www.anchor-lang.com/docs/updates/release-notes/0-32-1'),
]
for label, url in sources: link(label, url)
table(['Файл в coursework/rk1/evidence', 'Содержимое'], [
 ('localnet-run.json','9 транзакций, аккаунты, логи, balances, Counter, Profile, ELAB'),
 ('localnet-transaction-analysis.json','Разбор программ, инструкций и изменения lamports'),
 ('devnet-detective.json','Публичный сторонний пример и ссылка Explorer'),
 ('localnet-deploy-*.json','Program ID и реальные signatures деплоя'),
 ('mvp-integration.json','23 успешные проверки и signatures mint/check-in'),
 ('*-tests.txt и rust-cli.txt','Логи тестов и результат выполнения CLI'),
], [3.4,3.4])
p('Полные signatures сохранены в JSON без сокращений. Localnet genesis hash:')
code(lab['genesis_hash'])
p('Граница подтверждённых результатов: локальные программы и MVP работают; публичный '
  'Devnet-пример разобран; собственный Devnet-деплой остаётся следующим обязательным действием.')

doc.save(OUT / 'Evently_РК1_отчёт.docx')
(ROOT / 'docs/report.md').write_text('\n'.join(md), encoding='utf-8')
print(OUT / 'Evently_РК1_отчёт.docx')
